"""Converte a planilha oficial de classificação de eventos da CAPES (área de Computação) em CSV.

Uso:
    uv run --group qualis scripts/import_qualis.py                 # baixa a planilha 2017-2020 da CAPES
    uv run --group qualis scripts/import_qualis.py arquivo.xlsx 2021-2024   # outra planilha/período

A planilha deve ter uma aba com o cabeçalho "Sigla | Nome do evento | Estrato".
"""
import csv
import io
import sys
import urllib.request
from pathlib import Path

import openpyxl

URL_2017_2020 = (
    "https://www.gov.br/capes/pt-br/centrais-de-conteudo/documentos/avaliacao/"
    "09012022_CLASSIFICACAODEEVENTOSPARA20172020.xlsx/@@download/file"
)
OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "qualis"
ESTRATOS = {"A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4"}

# Erros conhecidos na planilha oficial (o estrato está correto, o nome não)
ERRATA_NOMES = {
    "CHI": "ACM Conference on Human Factors in Computing Systems",
}


def carregar_planilha(origem: str) -> openpyxl.Workbook:
    if origem.startswith("http"):
        with urllib.request.urlopen(origem, timeout=60) as resp:
            return openpyxl.load_workbook(io.BytesIO(resp.read()), read_only=True)
    return openpyxl.load_workbook(origem, read_only=True)


def extrair_eventos(wb: openpyxl.Workbook) -> list[tuple[str, str, str]]:
    for ws in wb.worksheets:
        linhas = ws.iter_rows(values_only=True)
        for linha in linhas:
            if linha and str(linha[0]).strip().lower() == "sigla":
                break
        else:
            continue
        eventos = []
        for sigla, nome, estrato, *_ in linhas:
            if not sigla or str(estrato).strip() not in ESTRATOS:
                continue
            sigla = " ".join(str(sigla).split())
            nome = ERRATA_NOMES.get(sigla, " ".join(str(nome or "").split()))
            eventos.append((sigla, nome, str(estrato).strip()))
        return eventos
    raise ValueError("Nenhuma aba com cabeçalho 'Sigla' encontrada")


def main() -> None:
    origem = sys.argv[1] if len(sys.argv) > 1 else URL_2017_2020
    periodo = sys.argv[2] if len(sys.argv) > 2 else "2017-2020"

    eventos = extrair_eventos(carregar_planilha(origem))
    saida = OUT_DIR / f"eventos_{periodo.replace('-', '_')}.csv"
    saida.parent.mkdir(parents=True, exist_ok=True)
    with open(saida, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["sigla", "nome", "estrato"])
        w.writerows(sorted(eventos, key=lambda e: e[0].lower()))
    print(f"{len(eventos)} eventos salvos em {saida}")


if __name__ == "__main__":
    main()
